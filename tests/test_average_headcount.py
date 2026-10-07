import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document

from average_headcount import (
    build_segments,
    choose_match,
    find_matches,
    Match,
    institution_regex,
    parse_ru_date,
    read_records,
    read_records_from_file,
)


def make_docx(path: Path, rows):
    document = Document()
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row in enumerate(rows):
        for cell_index, value in enumerate(row):
            table.cell(row_index, cell_index).text = value
    document.save(path)


def make_month_docx(path: Path, day_rows):
    document = Document()
    for report_date, rows in day_rows:
        document.add_paragraph("Суточная оперативная сводка")
        document.add_paragraph(f"На {report_date} года")
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_index, row in enumerate(rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value
        document.add_table(rows=1, cols=4).cell(0, 0).text = "Учреждение"
    document.save(path)


class AverageHeadcountTests(unittest.TestCase):
    def test_parse_ru_date_accepts_common_input_forms(self):
        expected = date(2026, 1, 20)
        self.assertEqual(parse_ru_date("20.01.2026"), expected)
        self.assertEqual(parse_ru_date("20-01-2026"), expected)
        self.assertEqual(parse_ru_date("20012026"), expected)
        self.assertEqual(parse_ru_date("200126"), expected)

    def test_institution_regex_does_not_match_prefix_numbers(self):
        pattern = institution_regex("ИК-2")
        self.assertIsNotNone(pattern.search("ИК-2- (1341)"))
        self.assertIsNone(pattern.search("ИК-25 бс- (758)"))

    def test_ic_does_not_match_ik_or_uic(self):
        pattern = institution_regex("ИЦ-1")
        self.assertIsNotNone(pattern.search("ИЦ-1 (115)"))
        self.assertIsNone(pattern.search("ИК-1- (1557)"))
        self.assertIsNone(pattern.search("УИЦ-1 (200)"))

    def test_find_matches_reads_nearest_number_to_the_right(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "01.01.2026.docx"
            make_docx(
                path,
                [
                    ["Строгий", "ИК-10- (1647) - Н", "ИК-10- (1647) - Н", "765", "765"],
                    ["Общий", "ИК-8- (1070) - Н", "ИК-8- (1070) - Н", "284", "284"],
                ],
            )

            matches = find_matches(path, "ИК-10")

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0].value, 765)
        self.assertEqual(matches[0].category, "Строгий")

    def test_carry_forward_segments_use_next_report_date(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            make_docx(root / "01.01.2026.docx", [["Строгий", "ИК-10", "", "100"]])
            make_docx(root / "01.02.2026.docx", [["Строгий", "ИК-10", "", "200"]])

            records = read_records(
                [root],
                "ИК-10",
                date(2026, 1, 1),
                date(2026, 2, 28),
                category=None,
                sum_matches=False,
                scan_all_tables=False,
            )
            segments = build_segments(records, date(2026, 1, 1), date(2026, 2, 28))

        self.assertEqual([(s.value, s.days, s.weighted) for s in segments], [(100, 31, 3100), (200, 28, 5600)])

    def test_reads_multiple_daily_reports_inside_one_month_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "01.01.2026.docx"
            make_month_docx(
                path,
                [
                    ("01.01.2026", [["Строгий", "ИК-10", "", "100"]]),
                    ("02.01.2026", [["Строгий", "ИК-10", "", "110"]]),
                ],
            )

            records = read_records_from_file(
                path,
                "ИК-10",
                category=None,
                sum_matches=False,
                scan_all_tables=False,
            )

        self.assertEqual([(r.report_date.isoformat(), r.value) for r in records], [("2026-01-01", 100), ("2026-01-02", 110)])

    def test_choose_match_prefers_main_institution_over_settlement(self):
        matches = [
            Match(857, "Строгий", "ИК-10- (1647) - Н", 0, 0, 3),
            Match(33, "Колония - поселение", "ИК-10- (52) -", 0, 1, 3),
        ]

        selected = choose_match(matches, "ИК-10")

        self.assertIsNotNone(selected)
        self.assertEqual(selected.value, 857)

    def test_choose_match_prefers_cell_that_starts_with_query(self):
        matches = [
            Match(947, "Строгий", "ИК-1- (1557) - В", 0, 0, 3),
            Match(28, "Принудительные работы", "УФИЦ ИК-1 (50)", 0, 1, 4),
        ]

        selected = choose_match(matches, "ИК-1")

        self.assertIsNotNone(selected)
        self.assertEqual(selected.value, 947)


if __name__ == "__main__":
    unittest.main()

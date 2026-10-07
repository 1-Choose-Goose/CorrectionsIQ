from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence


DATE_RE = re.compile(r"(?<!\d)(\d{2})\.(\d{2})\.(\d{4})(?!\d)")
REPORT_DATE_RE = re.compile(r"^\s*На\s+(\d{2}\.\d{2}\.\d{4})\s+года\s*$", re.IGNORECASE)
MAX_WORKERS = max(1, min(8, (os.cpu_count() or 2)))


@dataclass(frozen=True)
class Match:
    value: int
    category: str
    institution_cell: str
    table_index: int
    row_index: int
    value_cell_index: int


@dataclass(frozen=True)
class Record:
    report_date: date
    value: int
    path: Path
    matches: tuple[Match, ...]


@dataclass(frozen=True)
class Segment:
    report_date: date
    value: int
    days: int
    weighted: int
    path: Path
    matches: tuple[Match, ...]


def parse_ru_date(value: str) -> date:
    raw_value = value.strip()
    digits = re.sub(r"\D", "", raw_value)
    if len(digits) == 6:
        raw_value = f"{digits[:2]}.{digits[2:4]}.20{digits[4:]}"
    elif len(digits) == 8:
        raw_value = f"{digits[:2]}.{digits[2:4]}.{digits[4:]}"

    try:
        return datetime.strptime(raw_value, "%d.%m.%Y").date()
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Дата должна быть в формате ДД.ММ.ГГГГ, ДДММГГГГ или ДДММГГ: {value}"
        ) from exc


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).upper().replace("Ё", "Е")
    value = re.sub(r"[‐‑‒–—―]", "-", value)
    value = value.replace("«", " ").replace("»", " ").replace('"', " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def institution_regex(query: str) -> re.Pattern[str]:
    tokens = re.findall(r"[A-ZА-Я0-9]+", normalize_text(query))
    if not tokens:
        raise ValueError("Название учреждения пустое или не содержит букв/цифр.")
    pattern = r"(?<![A-ZА-Я0-9])" + r"[^A-ZА-Я0-9]*".join(map(re.escape, tokens))
    pattern += r"(?![A-ZА-Я0-9])"
    return re.compile(pattern)


def institution_starts_cell(query: str, cell_text: str) -> bool:
    query_tokens = re.findall(r"[A-ZА-Я0-9]+", normalize_text(query))
    cell_tokens = re.findall(r"[A-ZА-Я0-9]+", normalize_text(cell_text))
    return cell_tokens[: len(query_tokens)] == query_tokens


def is_secondary_category(category: str) -> bool:
    normalized = normalize_text(category)
    return any(
        marker in normalized
        for marker in (
            "КОЛОНИЯ - ПОСЕЛЕНИЕ",
            "ЕПКТ",
        )
    )


def choose_match(matches: Sequence[Match], institution: str) -> Match | None:
    candidates = [
        match for match in matches if institution_starts_cell(institution, match.institution_cell)
    ]
    if not candidates:
        candidates = list(matches)

    primary = [match for match in candidates if not is_secondary_category(match.category)]
    if len(primary) == 1:
        return primary[0]
    if primary:
        candidates = primary

    distinct = {(m.category, m.institution_cell, m.value) for m in candidates}
    if len(distinct) == 1:
        return candidates[0]
    return None


def parse_report_date_from_filename(path: Path) -> date | None:
    match = DATE_RE.search(path.stem)
    if match:
        return parse_ru_date(match.group(0))
    return None


def parse_report_date_from_document(document: Any) -> date | None:
    for paragraph in document.paragraphs[:30]:
        match = DATE_RE.search(paragraph.text)
        if match:
            return parse_ru_date(match.group(0))
    return None


def parse_report_date(path: Path) -> date | None:
    filename_date = parse_report_date_from_filename(path)
    if filename_date:
        return filename_date

    try:
        from docx import Document

        document = Document(path)
    except Exception:
        return None

    return parse_report_date_from_document(document)


def paragraph_report_date(text: str) -> date | None:
    match = REPORT_DATE_RE.search(text)
    if match:
        return parse_ru_date(match.group(1))
    return None


def iter_blocks(document: Any) -> Iterable[Any]:
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def is_standalone_int(value: str) -> bool:
    return bool(re.fullmatch(r"\s*\d[\d ]*\s*", value))


def to_int(value: str) -> int:
    return int(re.sub(r"\D", "", value))


def candidate_value(cells: Sequence[str], start_index: int) -> tuple[int, int] | None:
    for index in range(start_index + 1, len(cells)):
        if is_standalone_int(cells[index]):
            return to_int(cells[index]), index
    return None


def find_matches_in_table(
    table: Any,
    institution: str,
    category: str | None = None,
    table_index: int = 0,
) -> list[Match]:
    pattern = institution_regex(institution)
    category_filter = normalize_text(category) if category else None
    results: list[Match] = []
    seen: set[tuple[int, int, str, int]] = set()

    for row_index, row in enumerate(table.rows):
        cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
        row_category = cells[0] if cells else ""
        if category_filter and category_filter not in normalize_text(row_category):
            continue

        for cell_index, cell_text in enumerate(cells):
            # In auxiliary daily tables the institution is usually in the first
            # column and the numbers to the right are disciplinary indicators,
            # not the factual headcount.
            if cell_index == 0:
                continue
            if not pattern.search(normalize_text(cell_text)):
                continue
            found_value = candidate_value(cells, cell_index)
            if not found_value:
                continue
            value, value_cell_index = found_value
            key = (table_index, row_index, normalize_text(cell_text), value)
            if key in seen:
                continue
            seen.add(key)
            results.append(
                Match(
                    value=value,
                    category=row_category,
                    institution_cell=cell_text,
                    table_index=table_index,
                    row_index=row_index,
                    value_cell_index=value_cell_index,
                )
            )
    return results


def find_matches(
    path: Path,
    institution: str,
    category: str | None = None,
    scan_all_tables: bool = False,
) -> list[Match]:
    from docx import Document

    document = Document(path)
    tables = document.tables if scan_all_tables else document.tables[:1]
    results: list[Match] = []
    for table_index, table in enumerate(tables):
        results.extend(find_matches_in_table(table, institution, category, table_index))
    return results


def read_records_from_file(
    path: Path,
    institution: str,
    category: str | None,
    sum_matches: bool,
    scan_all_tables: bool,
) -> list[Record]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(path)
    except Exception as exc:
        print(f"Предупреждение: файл пропущен ({path}): {exc}", file=sys.stderr)
        return []
    fallback_date = parse_report_date_from_filename(path) or parse_report_date_from_document(document)
    current_date: date | None = None
    table_index = 0
    records: list[Record] = []

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            report_date = paragraph_report_date(block.text)
            if report_date:
                if fallback_date and (
                    report_date.year != fallback_date.year
                    or report_date.month != fallback_date.month
                ):
                    continue
                current_date = report_date
            continue

        if not isinstance(block, Table):
            continue

        table_date = current_date or fallback_date
        if table_date is None:
            table_index += 1
            continue

        matches = find_matches_in_table(block, institution, category, table_index)
        table_index += 1
        if not matches:
            continue

        if records and records[-1].report_date == table_date and not scan_all_tables:
            continue

        selected_match = None if sum_matches else choose_match(matches, institution)
        distinct = {(m.category, m.institution_cell, m.value) for m in matches}
        if selected_match is None and len(distinct) > 1 and not sum_matches:
            variants = "\n".join(
                f"  - категория: {m.category or '(пусто)'}, ячейка: {m.institution_cell}, "
                f"значение: {m.value}"
                for m in matches
            )
            raise RuntimeError(
                f"В файле {path} за {table_date:%d.%m.%Y} учреждение найдено неоднозначно:\n"
                f"{variants}\nУточните название учреждения."
            )

        value = sum(m.value for m in matches) if sum_matches else selected_match.value
        record_matches = tuple(matches if sum_matches else (selected_match,))
        records.append(Record(table_date, value, path, record_matches))

    return records


def iter_docx_files(roots: Iterable[Path]) -> Iterable[Path]:
    for root in roots:
        if root.is_file() and root.suffix.lower() == ".docx" and not root.name.startswith("~$"):
            yield root
        elif root.is_dir():
            for path in sorted(root.rglob("*.docx")):
                if not path.name.startswith("~$"):
                    yield path


def read_records(
    roots: Sequence[Path],
    institution: str,
    start: date,
    end: date,
    category: str | None,
    sum_matches: bool,
    scan_all_tables: bool,
) -> list[Record]:
    first_needed_file_date = start.replace(day=1)
    paths: list[Path] = []
    for path in iter_docx_files(roots):
        report_date = parse_report_date_from_filename(path)
        if report_date is not None and report_date > end:
            continue
        if report_date is not None and report_date < first_needed_file_date:
            continue
        paths.append(path)

    records: list[Record] = []
    if not paths:
        return []

    worker_count = min(MAX_WORKERS, len(paths))
    if worker_count == 1:
        for path in paths:
            file_records = read_records_from_file(
                path, institution, category, sum_matches, scan_all_tables
            )
            records.extend(record for record in file_records if record.report_date <= end)
    else:
        with ThreadPoolExecutor(max_workers=worker_count) as executor:
            futures = {
                executor.submit(
                    read_records_from_file,
                    path,
                    institution,
                    category,
                    sum_matches,
                    scan_all_tables,
                ): path
                for path in paths
            }
            for future in as_completed(futures):
                file_records = future.result()
                records.extend(record for record in file_records if record.report_date <= end)

    records.sort(key=lambda item: (item.report_date, str(item.path)))
    return collapse_same_date(records, sum_matches)


def collapse_same_date(records: Sequence[Record], sum_matches: bool) -> list[Record]:
    by_date: dict[date, list[Record]] = {}
    for record in records:
        by_date.setdefault(record.report_date, []).append(record)

    collapsed: list[Record] = []
    for report_date in sorted(by_date):
        same_day = by_date[report_date]
        if len(same_day) == 1:
            collapsed.append(same_day[0])
            continue
        if not sum_matches:
            files = "\n".join(f"  - {record.path}" for record in same_day)
            raise RuntimeError(
                f"Найдено несколько сводок за {report_date:%d.%m.%Y}:\n{files}\n"
                "Оставьте один файл или используйте --sum-matches, если их нужно сложить."
            )
        collapsed.append(
            Record(
                report_date=report_date,
                value=sum(record.value for record in same_day),
                path=same_day[0].path,
                matches=tuple(match for record in same_day for match in record.matches),
            )
        )
    return collapsed


def build_segments(records: Sequence[Record], start: date, end: date) -> list[Segment]:
    usable = [record for record in records if record.report_date <= end]
    if not usable:
        raise RuntimeError("Не найдено ни одной подходящей сводки.")

    active_index = None
    for index, record in enumerate(usable):
        if record.report_date <= start:
            active_index = index
        elif active_index is None:
            break

    if active_index is None:
        first_date = usable[0].report_date
        raise RuntimeError(
            f"Первая найденная сводка датирована {first_date:%d.%m.%Y}, "
            f"а период начинается {start:%d.%m.%Y}. Нет значения для начала периода."
        )

    segments: list[Segment] = []
    relevant = usable[active_index:]
    for index, record in enumerate(relevant):
        segment_start = max(record.report_date, start)
        next_date = relevant[index + 1].report_date if index + 1 < len(relevant) else end + timedelta(days=1)
        segment_end = min(end + timedelta(days=1), next_date)
        days = (segment_end - segment_start).days
        if days <= 0:
            continue
        segments.append(
            Segment(
                report_date=record.report_date,
                value=record.value,
                days=days,
                weighted=record.value * days,
                path=record.path,
                matches=record.matches,
            )
        )
    return segments


def write_csv(path: Path, segments: Sequence[Segment]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(["Дата сводки", "Значение", "Дней", "Взвешенная сумма", "Файл"])
        for segment in segments:
            writer.writerow(
                [
                    segment.report_date.strftime("%d.%m.%Y"),
                    segment.value,
                    segment.days,
                    segment.weighted,
                    str(segment.path),
                ]
            )


def write_xlsx(path: Path, segments: Sequence[Segment], average: float | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расчет"

    headers = ["Дата", "Количество"]
    sheet.append(headers)

    for segment in segments:
        sheet.append([segment.report_date, segment.value])

    total_days = sum(segment.days for segment in segments)
    total_weighted = sum(segment.weighted for segment in segments)
    calculated_average = average if average is not None else (
        total_weighted / total_days if total_days else 0
    )

    total_row = sheet.max_row + 2
    sheet.cell(total_row, 1, "Дней в периоде")
    sheet.cell(total_row, 2, total_days)
    sheet.cell(total_row + 1, 1, "Среднесписочное")
    sheet.cell(total_row + 1, 2, calculated_average)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="1F2937")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    for row_index in range(2, total_row):
        sheet.cell(row_index, 1).number_format = "dd.mm.yyyy"
        for col_index in (2,):
            sheet.cell(row_index, col_index).number_format = "#,##0"

    for row_index in (total_row, total_row + 1):
        sheet.cell(row_index, 1).font = Font(bold=True)
        sheet.cell(row_index, 2).font = Font(bold=True)
    sheet.cell(total_row + 1, 2).number_format = "#,##0.00"

    widths = [14, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def prompt_missing(args: argparse.Namespace) -> argparse.Namespace:
    if not args.root:
        raw = input("Папка со сводками [.]: ").strip() or "."
        args.root = [raw]
    if not args.institution:
        args.institution = input("Учреждение, например ИК-10: ").strip()
    if not args.start:
        args.start = parse_ru_date(input("Дата начала (ДД.ММ.ГГГГ): ").strip())
    if not args.end:
        args.end = parse_ru_date(input("Дата окончания (ДД.ММ.ГГГГ): ").strip())
    return args


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="CorrectionsIQ"
    )
    parser.add_argument("--root", action="append", help="Папка или файл .docx. Можно указать несколько раз.")
    parser.add_argument("--institution", help='Учреждение, например "ИК-10" или "СИЗО-1".')
    parser.add_argument("--category", help='Фильтр по первой колонке, например "Строгий".')
    parser.add_argument("--start", type=parse_ru_date, help="Начало периода: ДД.ММ.ГГГГ.")
    parser.add_argument("--end", type=parse_ru_date, help="Конец периода: ДД.ММ.ГГГГ.")
    parser.add_argument("--csv", type=Path, help="Путь для сохранения CSV-отчета.")
    parser.add_argument("--xlsx", type=Path, help="Путь для сохранения Excel-отчета.")
    parser.add_argument("--sum-matches", action="store_true", help="Сложить все найденные совпадения.")
    parser.add_argument("--scan-all-tables", action="store_true", help="Искать не только в первой таблице.")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = prompt_missing(parser.parse_args(argv))

    if not args.institution:
        parser.error("Укажите учреждение.")
    if args.start > args.end:
        parser.error("Дата начала не может быть позже даты окончания.")

    roots = [Path(root) for root in args.root]
    missing_roots = [str(root) for root in roots if not root.exists()]
    if missing_roots:
        parser.error("Не найдены пути: " + ", ".join(missing_roots))

    try:
        records = read_records(
            roots,
            args.institution,
            args.start,
            args.end,
            args.category,
            args.sum_matches,
            args.scan_all_tables,
        )
        segments = build_segments(records, args.start, args.end)
    except Exception as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1

    total_days = (args.end - args.start).days + 1
    total_weighted = sum(segment.weighted for segment in segments)
    average = total_weighted / total_days

    print(f"Учреждение: {args.institution}")
    if args.category:
        print(f"Категория: {args.category}")
    print(f"Период: {args.start:%d.%m.%Y} - {args.end:%d.%m.%Y} ({total_days} дн.)")
    print()
    print("Использованные значения:")
    for segment in segments:
        print(
            f"  {segment.report_date:%d.%m.%Y}: {segment.value} x {segment.days} дн. "
            f"= {segment.weighted} ({segment.path})"
        )
    print()
    print(f"Сумма человеко-дней: {total_weighted}")
    print(f"Среднесписочное количество: {average:.2f}")

    covered_days = sum(segment.days for segment in segments)
    if covered_days != total_days:
        print(f"Предупреждение: покрыто {covered_days} из {total_days} дней.", file=sys.stderr)

    if args.csv:
        write_csv(args.csv, segments)
        print(f"CSV сохранен: {args.csv}")
    if args.xlsx:
        write_xlsx(args.xlsx, segments, average)
        print(f"Excel сохранен: {args.xlsx}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
